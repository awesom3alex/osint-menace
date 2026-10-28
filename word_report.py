"""Generate editable Word reports from saved case data."""

from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from case_workflow import report_case
from dossier import SECTIONS, dossier_values
from readable_results import readable_text


def clean_text(value):
    text = str(value)
    return "".join(
        char
        for char in text
        if char in "\n\r\t"
        or 0x20 <= ord(char) <= 0xD7FF
        or 0xE000 <= ord(char) <= 0xFFFD
        or 0x10000 <= ord(char) <= 0x10FFFF
    )


def build_word_report(case):
    if case.get("report_builder") is not None or case.get("investigation", {}).get("assessments") or case.get("investigation", {}).get("redaction", {}).get("enabled"):
        from report_builder import word
        return word(case)
    case = report_case(case)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    for border in document.styles["Title"].element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.space_after = Pt(7)
    document.core_properties.title = clean_text(case["name"])
    document.core_properties.author = "OSINT Menace"

    document.add_paragraph("Investigation Dossier", "Title")
    document.add_paragraph(clean_text(case["name"]), "Subtitle")
    document.add_paragraph("OSINT Menace | A Defensive Thinking Project.")
    for label, key in (("Case ID", "id"), ("Created", "created_at"), ("Updated", "updated_at")):
        document.add_paragraph(label + ": " + clean_text(case.get(key, "")))
    if case.get("notes"):
        document.add_paragraph(clean_text(case["notes"]))

    values = dossier_values(case)
    for title, fields in SECTIONS:
        if not any(values[key] for key, _ in fields):
            continue
        document.add_heading(title, level=1)
        filled = False
        for key, label in fields:
            if values[key]:
                filled = True
                document.add_heading(label, level=2)
                document.add_paragraph(clean_text(values[key]))
        if not filled:
            document.add_paragraph("Not recorded.")

    document.add_heading("Saved Evidence", level=1)
    for search in case.get("searches", []):
        document.add_heading(clean_text(search.get("query", "Saved record")), level=2)
        if search.get("review_status") or search.get("analyst_notes") or search.get("tags"):
            document.add_paragraph(
                "Review status: " + clean_text(search.get("review_status", "unreviewed"))
            )
            document.add_paragraph("Tags: " + clean_text(", ".join(search.get("tags", []))))
            if search.get("analyst_notes"):
                document.add_paragraph("Analyst notes: " + clean_text(search["analyst_notes"]))
        document.add_paragraph(
            "Record: "
            + clean_text(search.get("id", ""))
            + "\nSaved at: "
            + clean_text(search.get("created_at", ""))
            + "\nModule: "
            + clean_text(search.get("search_type", ""))
        )
        for result in search.get("results", []):
            document.add_heading(clean_text(result.get("name", "Result")), level=3)
            for finding in result.get("findings", []):
                if not finding.get("include_in_report", True):
                    continue
                paragraph = document.add_paragraph()
                paragraph.add_run(clean_text(finding.get("label", "Finding")) + ": ").bold = True
                value = finding.get("value", "")
                if isinstance(value, (dict, list)):
                    value = readable_text(value)
                paragraph.add_run(clean_text(value))
                for source in finding.get("merged_sources", []):
                    document.add_paragraph(
                        "Merged source: "
                        + clean_text(source["source"])
                        + " | "
                        + clean_text(source["collected_at"])
                        + " | "
                        + clean_text(source["verification_status"])
                        + "\n"
                        + clean_text(source["source_url"])
                        + "\n"
                        + clean_text(source["analyst_notes"])
                    )
                if finding.get("source"):
                    document.add_paragraph("Source: " + clean_text(finding["source"]))
            for link in result.get("links", []):
                document.add_paragraph(
                    clean_text(link.get("label", "Source"))
                    + "\n"
                    + clean_text(link.get("value", ""))
                )
            for note in result.get("notes", []):
                document.add_paragraph("Note: " + clean_text(note))
    if not case.get("searches"):
        document.add_paragraph("No saved evidence.")
    from graph_reports import attachments, legend_ink, report_details, sections, summary_lines

    for graph_section in sections(case):
        document.add_page_break()
        document.add_heading("Case graph " + clean_text(graph_section["graph"]["name"]), level=1)
        picture = document.add_picture(BytesIO(graph_section["image"]), width=Inches(6.7))
        picture._inline.docPr.set(
            "descr", "Investigation graph " + clean_text(graph_section["graph"]["name"])
        )
        for line in summary_lines(graph_section)[:2]:
            document.add_paragraph(clean_text(line))
        filters = graph_section["filters"]
        if filters.get("from") or filters.get("to"):
            document.add_paragraph(
                f"Date range: {filters.get('from') or 'Any date'} to {filters.get('to') or 'Any date'}. "
                + ("Undated connections included." if filters.get("includeUndated", True) else "Undated connections excluded."))
        document.add_heading("Graph legend", level=2)
        for entry in graph_section["graph"]["legend"]:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(" " + clean_text(entry["type"]) + " ")
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(legend_ink(entry["color"])[1:])
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), entry["color"].lstrip("#"))
            run._r.get_or_add_rPr().append(shading)
            paragraph.add_run("  " + clean_text(entry["description"]))
        for title_text, records in report_details(graph_section):
            if not records:
                continue
            document.add_heading(title_text, level=2)
            for label, value in records:
                document.add_heading(clean_text(label), level=3)
                document.add_paragraph(clean_text(value))

        for path, caption in attachments(graph_section, case["id"]):
            picture = document.add_picture(str(path), width=Inches(6.7))
            picture._inline.docPr.set("descr", clean_text(caption))
            document.paragraphs[-1].paragraph_format.keep_with_next = True
            document.add_paragraph("Supporting screenshot: " + clean_text(caption))

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
