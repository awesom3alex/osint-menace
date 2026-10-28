"""Views, durable milestones, and report selection regression coverage."""

from unittest import TestCase, main

from docx import Document

import app
import case_backups
import graph_reports
import pdf_reports
from tests import test_case_graphs as base
from word_report import build_word_report


class MilestoneTests(TestCase):
    setUp = base.GraphTests.setUp
    put = base.GraphTests.put
    populated = base.GraphTests.populated

    def view(self, graph):
        return dict(
            id="a" * 32,
            name="People",
            filters=dict(
                query="",
                type="Person",
                status="",
                **{"from": "2026-10-01", "to": "2026-10-31"},
                includeUndated=False,
            ),
            zoom=1,
            pan=dict(x=0, y=0),
            positions={n["id"]: n["position"] for n in graph["nodes"]},
            collapsed=[],
            groupStates={},
        )

    def snapshot(self, graph, name="Milestone"):
        return self.client.post(
            self.base + "/" + graph["id"] + "/snapshots",
            json=dict(name=name, revision=graph["revision"]),
        )

    def test_restore_keeps_pre_restore_graph_and_survives_backup(self):
        graph = self.populated()
        graph["views"] = [self.view(graph)]
        graph["report"] = dict(include=True, viewId="a" * 32)
        self.assertEqual(self.put(graph).status_code, 200)
        graph = app.load_case(self.case["id"])["graphs"][0]
        created = self.snapshot(graph)
        self.assertEqual(created.status_code, 200, created.data)
        graph = created.json["graph"]
        sid = graph["snapshots"][0]["id"]
        graph["nodes"][0]["label"] = "Changed after milestone"
        self.assertEqual(self.put(graph).status_code, 200)
        graph = app.load_case(self.case["id"])["graphs"][0]
        restored = self.client.post(
            self.base + "/" + graph["id"] + "/snapshots/" + sid + "/restore",
            json={"revision": graph["revision"]},
        )
        self.assertEqual(restored.status_code, 200, restored.data)
        graph = restored.json["graph"]
        self.assertEqual(graph["nodes"][0]["label"], "Alice")
        self.assertEqual(len(graph["snapshots"]), 2)
        self.assertEqual(
            graph["snapshots"][1]["data"]["nodes"][0]["label"], "Changed after milestone"
        )
        self.assertEqual(graph["views"][0]["name"], "People")
        with case_backups.archive_case(app.load_case(self.case["id"])) as archive:
            path = app.CASES_DIR.parent / "graph.zip"
            path.write_bytes(archive.read())
        restored_case, _, _ = case_backups.checked_zip(path)
        self.assertEqual(restored_case["graphs"][0], graph)

    def test_stale_snapshot_and_client_history_replacement_are_rejected(self):
        created = self.snapshot(self.graph)
        graph = created.json["graph"]
        self.assertEqual(self.snapshot(self.graph).status_code, 409)
        graph["snapshots"] = []
        self.assertEqual(self.put(graph).status_code, 400)
        self.assertEqual(len(app.load_case(self.case["id"])["graphs"][0]["snapshots"]), 1)

    def test_snapshot_limit_never_discards_existing_milestones(self):
        graph = self.graph
        for i in range(10):
            result = self.snapshot(graph, f"Milestone {i}")
            self.assertEqual(result.status_code, 200)
            graph = result.json["graph"]
        before = app.case_path(self.case["id"]).read_bytes()
        self.assertEqual(self.snapshot(graph, "Overflow").status_code, 400)
        sid = graph["snapshots"][0]["id"]
        response = self.client.post(
            self.base + "/" + graph["id"] + "/snapshots/" + sid + "/restore",
            json={"revision": graph["revision"]},
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(app.case_path(self.case["id"]).read_bytes(), before)

    def test_invalid_views_and_nested_snapshots_are_rejected(self):
        for field, value in [
            ("zoom", float("inf")),
            ("positions", {"bad": {"x": "NaN", "y": 1}}),
            ("filters", {"from": "2026-11-01", "to": "2026-10-01"}),
        ]:
            graph = self.populated()
            view = self.view(graph)
            view[field] = value
            graph["views"] = [view]
            self.assertEqual(self.put(graph).status_code, 400)
        graph = self.populated()
        graph["report"] = dict(include=True, viewId="missing")
        self.assertEqual(self.put(graph).status_code, 400)
        graph = self.snapshot(self.graph).json["graph"]
        graph["snapshots"][0]["data"]["snapshots"] = []
        with self.assertRaises(ValueError):
            case_backups.validate_case(dict(app.load_case(self.case["id"]), graphs=[graph]))

    def test_remove_snapshot_and_foreign_case_cannot_access_history(self):
        graph = self.snapshot(self.graph).json["graph"]
        sid = graph["snapshots"][0]["id"]
        other = app.create_case("Other case", "")
        self.assertEqual(
            self.client.post(
                f"/cases/{other['id']}/graphs/{graph['id']}/snapshots",
                json={"revision": graph["revision"]},
            ).status_code,
            404,
        )
        result = self.client.delete(
            self.base + "/" + graph["id"] + "/snapshots/" + sid,
            json={"revision": graph["revision"]},
        )
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json["graph"]["snapshots"], [])

    def report_case(self):
        graph = self.populated()
        graph["nodes"][0]["photo"] = ""
        graph["nodes"][1]["type"] = "Account"
        graph["edges"][0].update(date="2026-10-27", notes="Relationship evidence", evidenceRefs=[])
        graph["views"] = [self.view(graph)]
        graph["views"][0]["filters"]["type"] = ""
        graph["report"] = dict(include=True, viewId="a" * 32)
        return dict(self.case, graphs=[graph])

    def test_report_view_date_selection_and_both_formats_embed_figure(self):
        case = self.report_case()
        section = graph_reports.sections(case)[0]
        self.assertEqual(len(section["graph"]["edges"]), 1)
        self.assertTrue(section["image"].startswith(b"\x89PNG"))
        word = Document(build_word_report(case))
        self.assertEqual(len(word.inline_shapes), 1)
        text = "\n".join(p.text for p in word.paragraphs)
        self.assertIn("View: People", text)
        self.assertIn("Relationship evidence", text)
        pdf = pdf_reports.generate(case)
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertIn(b"/Subtype /Image", pdf)
        case["graphs"][0]["views"][0]["filters"]["to"] = "2026-10-20"
        self.assertEqual(graph_reports.sections(case)[0]["graph"]["edges"], [])
        case["graphs"][0]["report"]["include"] = False
        self.assertEqual(graph_reports.sections(case), [])

    def test_supporting_screenshots_are_included_in_word_and_pdf(self):
        from unittest.mock import patch

        from PIL import Image

        import case_features

        case = self.report_case()
        eid = "e" * 32
        finding = dict(
            label="Screenshot source",
            value="Linked evidence",
            source="Fixture",
            evidence=[dict(id=eid, caption="Supporting capture")],
        )
        section = dict(name="Fixture", findings=[finding], links=[], notes=[])
        search = app.make_search_record(
            "username", "sample", [section], app.result_summary([section])
        )
        case["searches"] = [search]
        case["graphs"][0]["edges"][0]["evidenceRefs"] = [search["id"] + ":0:0"]
        folder = app.CASES_DIR.parent / "screenshots"
        path = folder / case["id"] / (eid + ".png")
        path.parent.mkdir(parents=True)
        Image.new("RGB", (40, 20), "red").save(path)
        with patch.object(case_features, "EVIDENCE_DIR", folder):
            word = Document(build_word_report(case))
            self.assertEqual(len(word.inline_shapes), 2)
            self.assertIn(
                "Supporting screenshot: Supporting capture",
                "\n".join(p.text for p in word.paragraphs),
            )
            self.assertTrue(pdf_reports.generate(case).startswith(b"%PDF"))

    def test_excluded_finding_is_not_leaked_through_graph(self):
        case = self.report_case()
        section = dict(
            name="Private",
            findings=[
                dict(
                    label="Excluded source", value="SECRET_NOT_FOR_REPORT", include_in_report=False
                )
            ],
            links=[],
            notes=[],
        )
        search = app.make_search_record(
            "username", "example", [section], app.result_summary([section])
        )
        case["searches"] = [search]
        ref = search["id"] + ":0:0"
        case["graphs"][0]["nodes"][0]["source"] = ref
        case["graphs"][0]["nodes"][0]["label"] = "SECRET_NOT_FOR_REPORT"
        output = graph_reports.sections(case)[0]
        self.assertFalse(any(n["id"] == "person" for n in output["graph"]["nodes"]))
        self.assertEqual(output["graph"]["edges"], [])
        self.assertNotIn("SECRET_NOT_FOR_REPORT", "\n".join(graph_reports.summary_lines(output)))


if __name__ == "__main__":
    main()
