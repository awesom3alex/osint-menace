"""Focused coverage of report composition, persistence and export consistency."""
import copy
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from docx import Document
from PIL import Image

import app
import case_backups
import case_features
import pdf_reports
import report_builder as builder
from word_report import build_word_report


class ReportBuilderTests(TestCase):
    def setUp(self):
        folder = TemporaryDirectory(); self.addCleanup(folder.cleanup)
        root = Path(folder.name)
        for obj, name, value in [(app, 'CASES_DIR', root / 'cases'), (case_features, 'EVIDENCE_DIR', root / 'evidence')]:
            patcher = patch.object(obj, name, value); patcher.start(); self.addCleanup(patcher.stop)
        self.case = app.create_case('Reader report', 'Context notes')
        self.case['searches'] = [dict(id='search', query='alice', search_type='username', created_at='2026-10-27',
                                     results=[dict(name='Accounts', findings=[dict(label='Public profile', value='alice.example', source='Example', verification_status='Unverified', evidence=[dict(id='e'*32, caption='Profile capture')]), dict(label='Excluded fact', value='DO NOT EXPORT', include_in_report=False)])])]
        app.save_case(self.case)
        directory = case_features.EVIDENCE_DIR / self.case['id']; directory.mkdir(parents=True)
        Image.new('RGB', (200, 80), 'white').save(directory / ('e'*32 + '.png'))
        self.client = app.app.test_client(); self.url = '/cases/' + self.case['id']

    def save(self, selected=None):
        config = builder.settings(self.case)
        config['summary'] = dict(purpose='Assess the public profile.', findings='One profile located.', questions='Ownership remains unverified.')
        if selected is not None:
            config['selected'] = selected
        config['order'].reverse()
        return self.client.post(self.url + '/report-builder', json=dict(config, updated_at=self.case['updated_at']))

    def test_persistence_preview_and_exports_share_order_and_numbers(self):
        self.assertEqual(self.save().status_code, 200)
        case = app.load_case(self.case['id']); data = builder.compose(case)
        self.assertEqual(data['blocks'][0]['kind'], 'screenshot')
        self.assertIn('Evidence 1', data['blocks'][0]['label'])
        self.assertIn('Evidence 1', data['appendix'][0]['label'])
        self.assertFalse(data['warnings'])
        for suffix in ['/report-preview', '/report.html']:
            response = self.client.get(self.url + suffix)
            self.assertEqual(response.status_code, 200)
            self.assertIn(b'Assess the public profile.', response.data)
            self.assertNotIn(b'DO NOT EXPORT', response.data)
        self.assertTrue(pdf_reports.generate(case).startswith(b'%PDF'))
        doc = Document(build_word_report(case)); text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Evidence sources', text); self.assertNotIn('DO NOT EXPORT', text)
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertIn('Evidence 1', app.render_markdown_report(case))
        self.assertEqual(self.client.get(self.url + '/report-preview.pdf').headers['Content-Disposition'], 'inline; filename="report-preview.pdf"')

    def test_selection_does_not_modify_originals_and_new_content_is_unselected(self):
        self.assertEqual(self.save(['notes']).status_code, 200)
        case = app.load_case(self.case['id'])
        self.assertEqual([b['kind'] for b in builder.compose(case)['blocks']], ['notes'])
        self.assertEqual(case['searches'], self.case['searches'])
        case['searches'][0]['results'][0]['findings'].append(dict(label='New', value='NEW'))
        self.assertNotIn('finding:search:0:2', builder.settings(case)['selected'])

    def test_conflicts_invalid_ids_and_backup_validation(self):
        self.assertEqual(self.save().status_code, 200)
        self.assertEqual(self.save().status_code, 409)
        case = app.load_case(self.case['id']); config = builder.settings(case)
        config['selected'].append('not-a-real-item')
        response = self.client.post(self.url + '/report-builder', json=dict(config, updated_at=case['updated_at']))
        self.assertEqual(response.status_code, 400)
        invalid = copy.deepcopy(case); invalid['report_builder']['summary']['purpose'] = []
        with self.assertRaises(ValueError):
            case_backups.validate_case(invalid)

    def test_empty_and_missing_summary_warnings_and_escaped_content(self):
        self.assertEqual(len(builder.compose(self.case)['warnings']), 3)
        self.assertEqual(self.save([]).status_code, 200)
        case = app.load_case(self.case['id'])
        self.assertIn('No report content selected.', builder.compose(case)['warnings'])
        case['report_builder']['summary']['purpose'] = '<script>alert(1)</script>'; app.save_case(case)
        response = self.client.get(self.url + '/report-preview')
        self.assertIn(b'&lt;script&gt;', response.data)
        self.assertNotIn(b'<script>alert', response.data)

    def test_graph_sources_numbered_without_printed_hex_and_exclusions_preserved(self):
        self.client.post(self.url + '/graphs', data=dict(name='Relationships'))
        case = app.load_case(self.case['id']); graph = case['graphs'][0]
        graph['nodes'] = [dict(id='a', label='Alice', type='Person', position=dict(x=0,y=0), fields=[], source='search:0:0'), dict(id='b', label='Hidden', type='Person', position=dict(x=100,y=0), fields=[], source='search:0:1')]
        graph.setdefault('report', {})['include'] = True
        config = builder.settings(case); case['report_builder'] = config
        data = builder.compose(case)
        block = next(b for b in data['blocks'] if b['kind'] == 'graph')
        self.assertIn('Evidence 1', block['text'])
        self.assertIn('Sources: Evidence 1', str(block['details']))
        self.assertNotIn('Hidden', str(block['details']))
        doc = Document(builder.word(case)); text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertNotRegex(text, r'#[a-fA-F0-9]{6}')
        self.assertTrue(doc.element.xpath('//w:rPr/w:shd'))
        self.assertTrue(builder.pdf(case).startswith(b'%PDF'))

    def test_cli_html_respects_composition_and_redaction(self):
        from menace import export_case
        self.case['investigation'] = dict(redaction=dict(enabled=True, terms=['alice', 'Reader']))
        self.case['report_builder'] = builder.settings(self.case)
        self.case['report_builder']['summary']['purpose'] = 'Review alice'
        with TemporaryDirectory() as folder:
            path = Path(folder) / 'report.html'
            export_case(self.case, 'html', path)
            html = path.read_text()
        self.assertNotIn('alice', html.lower())
        self.assertNotIn('Reader', html)
        self.assertNotIn('DO NOT EXPORT', html)
        self.assertNotIn('<iframe', html)
        self.assertIn('[REDACTED]', html)

    def test_graph_selection_syncs_both_directions_and_invalidates_old_revision(self):
        self.client.post(self.url + '/graphs', data=dict(name='Relationships'))
        self.case = app.load_case(self.case['id'])
        graph = self.case['graphs'][0]
        key = 'graph:' + graph['id']
        self.assertEqual(self.save([key]).status_code, 200)
        case = app.load_case(self.case['id'])
        current = case['graphs'][0]
        self.assertTrue(current['report']['include'])
        self.assertGreater(current['revision'], graph['revision'])
        endpoint = self.url + '/graphs/' + graph['id']
        self.assertEqual(self.client.put(endpoint, json=graph).status_code, 409)
        current['report']['include'] = False
        self.assertEqual(self.client.put(endpoint, json=current).status_code, 200)
        case = app.load_case(case['id'])
        self.assertNotIn(key, builder.settings(case)['selected'])
        self.assertFalse(any(b['kind'] == 'graph' for b in builder.compose(case)['blocks']))
        current = case['graphs'][0]
        current['report']['include'] = True
        self.assertEqual(self.client.put(endpoint, json=current).status_code, 200)
        self.assertIn(key, builder.settings(app.load_case(case['id']))['selected'])

    def test_merged_secondary_source_retains_provenance_and_primary_exclusion(self):
        self.client.post(self.url + '/graphs', data=dict(name='Relationships'))
        case = app.load_case(self.case['id'])
        secondary = case['searches'][0]['results'][0]['findings'][1]
        secondary.update(include_in_report=True, label='Secondary profile', value='public account', source='Secondary source')
        case['merge_groups'] = [dict(id='m', members=['search:0:0', 'search:0:1'])]
        graph = case['graphs'][0]
        graph['nodes'] = [dict(id='a', label='Secondary', type='Person', position=dict(x=0, y=0), fields=[], source='search:0:1')]
        graph['report'] = dict(include=True)
        data = builder.compose(case)
        self.assertIn('Secondary source', str(data['appendix']))
        self.assertNotIn('Source record unavailable', str(data['appendix']))
        self.assertEqual(sum(b['kind'] == 'finding' for b in data['blocks']), 1)
        case['searches'][0]['results'][0]['findings'][0]['include_in_report'] = False
        data = builder.compose(case)
        self.assertNotIn('Secondary source', str(data['appendix']))
        self.assertNotIn('Secondary', str(next(b for b in data['blocks'] if b['kind'] == 'graph')['details']))
