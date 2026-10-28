"""Ordered report composition shared by preview, PDF and Word."""
import base64
import copy
import re
from io import BytesIO

from flask import Blueprint, Response, abort, jsonify, render_template, request

from readable_results import readable_text

bp = Blueprint('report_builder', __name__)
SUMMARY_FIELDS = (('purpose', 'Investigation purpose'), ('findings', 'Key findings'), ('questions', 'Outstanding questions'))


def inventory(case):
    from case_workflow import findings, report_case
    from dossier import SECTIONS, dossier_values
    items = []
    if case.get('notes'):
        items.append(dict(id='notes', kind='notes', label='Case notes', text=case['notes']))
    values = dossier_values(case)
    for title, fields in SECTIONS:
        text = '\n\n'.join(label + ': ' + str(values[key]) for key, label in fields if values[key])
        if text:
            items.append(dict(id='dossier:' + title, kind='notes', label=title, text=text))
    for ref, search, section, finding in findings(report_case(case)):
        if not finding.get('include_in_report', True):
            continue
        label = finding.get('label', 'Finding')
        items.append(dict(id='finding:' + ref, kind='finding', label=label, ref=ref,
                          text=readable_text(finding.get('value', '')), finding=finding,
                          source=finding.get('source') or section.get('module', {}).get('name', 'OSINT Menace'),
                          collected=finding.get('collected_at') or search.get('created_at', '')))
        for evidence in finding.get('evidence', []):
            eid = evidence.get('id', '')
            if re.fullmatch(r'[a-f0-9]{32}', eid):
                items.append(dict(id='screenshot:' + ref + ':' + eid, kind='screenshot',
                                  label=evidence.get('caption') or label, ref=ref, eid=eid, text=''))
    for task in case.get('investigation', {}).get('tasks', []):
        items.append(dict(id='task:' + task['id'], kind='notes', label='Task: ' + task['title'], default=False,
                          text='Status: ' + task['status'] + ('\nDue: ' + task['due'] if task.get('due') else '') + '\n' + task.get('notes', '')))
    for entry in case.get('investigation', {}).get('timeline', []):
        items.append(dict(id='event:' + entry['id'], kind='notes', label='Timeline: ' + entry['title'], default=False,
                          text=entry['timestamp'] + '\n' + entry.get('notes', '')))
    for graph in case.get('graphs', []):
        items.append(dict(id='graph:' + graph['id'], kind='graph', label=graph['name'], text='',
                          default=bool(graph.get('report', {}).get('include'))))
    return items


def settings(case):
    available = inventory(case)
    saved = case.get('report_builder')
    if saved is None:
        return dict(summary={key: '' for key, _ in SUMMARY_FIELDS}, order=[i['id'] for i in available],
                    selected=[i['id'] for i in available if i.get('default', True)])
    ids = {i['id'] for i in available}
    order = [i for i in saved['order'] if i in ids]
    return dict(summary=saved['summary'], order=order + [i['id'] for i in available if i['id'] not in order],
                selected=[i for i in saved['selected'] if i in ids])


def compose(case):
    from graph_reports import report_details, sections
    config = settings(case)
    lookup = {i['id']: i for i in inventory(case)}
    selected = set(config['selected'])
    chosen = [lookup[k] for k in config['order'] if k in selected]
    graph_case = copy.deepcopy(case)
    for graph in graph_case.get('graphs', []):
        graph.setdefault('report', {})['include'] = 'graph:' + graph['id'] in selected
    graphs = {s['graph']['id']: s for s in sections(graph_case)}
    refs = []
    for item in chosen:
        if item.get('ref') and item['ref'] not in refs:
            refs.append(item['ref'])
        if item['kind'] == 'graph':
            for evidence in graphs[item['id'][6:]]['evidence']:
                if evidence['ref'] not in refs:
                    refs.append(evidence['ref'])
    from investigation import assessment_text, redact_report
    source_lookup = {i['ref']: i for i in lookup.values() if i['kind'] == 'finding'}
    from case_workflow import findings
    originals = {ref: (search, section, finding) for ref, search, section, finding in findings(case)}
    for group in case.get('merge_groups', []):
        members = [ref for ref in group['members'] if ref in originals]
        if not members or members[0] not in source_lookup:
            continue
        for ref in members[1:]:
            search, section, finding = originals[ref]
            if not finding.get('include_in_report', True):
                continue
            source_lookup[ref] = dict(ref=ref, label=finding.get('label', 'Finding'), finding=finding,
                                     text=readable_text(finding.get('value', '')),
                                     source=finding.get('source') or section.get('module', {}).get('name', 'OSINT Menace'),
                                     collected=finding.get('collected_at') or search.get('created_at', ''))
    allowed_refs = set(source_lookup)
    for item in chosen:
        target_keys = [item['ref']] if item['kind'] == 'finding' else []
        if item['kind'] == 'graph':
            section = graphs[item['id'][6:]]
            target_keys = ['edge:' + item['id'][6:] + ':' + edge['id'] for edge in section['graph']['edges']]
        for key in target_keys:
            for ref in case.get('investigation', {}).get('assessments', {}).get(key, {}).get('refs', []):
                if ref in allowed_refs and ref not in refs:
                    refs.append(ref)
    numbers = {ref: 'Evidence ' + str(index + 1) for index, ref in enumerate(refs)}
    appendix = []
    for ref in refs:
        source = source_lookup.get(ref)
        if source:
            f = source['finding']
            text = ('Source: ' + source['source'] + '\nCollected: ' + (source['collected'] or 'Not recorded')
                    + '\nAssessment: ' + f.get('verification_status', 'Unverified')
                    + (('\nCaptured information: ' + source['text']) if not case.get('is_demo') else ''))
            if f.get('source_url'):
                text += '\nSource page: ' + f['source_url']
            for merged in f.get('merged_sources', []):
                text += '\nAdditional source: ' + str(merged.get('source', '')) + ' | ' + str(merged.get('source_url', ''))
            appendix.append(dict(label=numbers[ref] + ' — ' + source['label'], text=text + '\nCase record: ' + ref))
        else:
            appendix.append(dict(label=numbers[ref], text='Source record unavailable. Case record: ' + ref))
    warnings = [label + ' has not been written.' for key, label in SUMMARY_FIELDS if not config['summary'].get(key, '').strip()]
    if not chosen:
        warnings.append('No report content selected.')
    blocks = []
    for item in chosen:
        block = dict(label=item['label'], text=item['text'], kind=item['kind'])
        if item.get('ref'):
            block['label'] = numbers[item['ref']] + ' — ' + block['label']
        if item['kind'] == 'finding':
            f = item['finding']
            block['text'] += assessment_text(case, item['ref'], numbers)
            block['text'] = 'Assessment: ' + f.get('verification_status', 'Unverified') + '\n' + block['text']
            if f.get('analyst_notes'):
                block['text'] += '\nAnalyst notes: ' + f['analyst_notes']
        if item['kind'] == 'screenshot':
            import case_features
            path = case_features.EVIDENCE_DIR / case['id'] / (item['eid'] + '.png')
            from demo import DEMO_ID
            if case['id'] == DEMO_ID:
                from pathlib import Path
                path = Path(__file__).parent / 'static' / 'demo' / (item['eid'] + '.png')
            if path.is_file():
                block['image'] = path.read_bytes()
            else:
                block['text'] = 'Screenshot file unavailable.'
                warnings.append(item['label'] + ': screenshot file unavailable.')
        if item['kind'] == 'graph':
            section = graphs[item['id'][6:]]
            block['image'] = section['image']
            block['legend'] = section['graph']['legend']
            block['text'] = 'View: ' + section['view'] + '\nSolid arrows indicate supported connections; dashed arrows indicate hypotheses.'
            filters = section['filters']
            if filters.get('from') or filters.get('to'):
                block['text'] += '\nDate range: ' + (filters.get('from') or 'Any date') + ' to ' + (filters.get('to') or 'Any date') + ('. Undated connections included.' if filters.get('includeUndated', True) else '. Undated connections excluded.')
            block['details'] = list(report_details(section))[:2]
            for records, entities in ((block['details'][0][1], section['graph']['edges']),
                                      (block['details'][1][1], section['graph']['nodes'])):
                for index, entity in enumerate(entities):
                    entity_refs = entity.get('evidenceRefs', []) if 'source' in entity and 'target' in entity else [entity.get('source', ''), *entity.get('sourceRefs', [])]
                    citations = list(dict.fromkeys(numbers[r] for r in entity_refs if r in numbers))
                    if citations:
                        title, value = records[index]
                        records[index] = (title, value + '\nSources: ' + ', '.join(citations))
                    if 'target' in entity:
                        title, value = records[index]
                        records[index] = (title, value + assessment_text(case, 'edge:' + item['id'][6:] + ':' + entity['id'], numbers))
            evidence_labels = [numbers[e['ref']] for e in section['evidence']]
            if evidence_labels:
                block['text'] += '\nLinked evidence: ' + ', '.join(evidence_labels)
            if not section['graph']['nodes']:
                warnings.append(item['label'] + ': selected graph view contains no visible objects.')
        if block.get('image'):
            block['image_url'] = 'data:image/png;base64,' + base64.b64encode(block['image']).decode()
        blocks.append(block)
    return redact_report(dict(case=case, summary=[(label, config['summary'].get(key, '')) for key, label in SUMMARY_FIELDS],
                blocks=blocks, appendix=appendix, warnings=warnings))


@bp.get('/cases/<case_id>/report-builder')
def editor(case_id):
    import app
    case = app.load_case(case_id)
    items = {i['id']: i for i in inventory(case)}
    config = settings(case)
    return render_template('report_builder.html', case=case, config=config,
                           items=[items[k] for k in config['order']], fields=SUMMARY_FIELDS)


@bp.post('/cases/<case_id>/report-builder')
def save(case_id):
    import app
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        abort(400)
    with app.CASE_LOCK:
        case = app.load_case(case_id)
        if data.get('updated_at') != case.get('updated_at'):
            return jsonify(error='The case changed. Reload the builder before saving.'), 409
        ids = {i['id'] for i in inventory(case)}
        order, selected, summary = data.get('order'), data.get('selected'), data.get('summary')
        if (not isinstance(order, list) or not all(isinstance(x, str) for x in order)
            or len(order) != len(set(order)) or set(order) != ids
            or not isinstance(selected, list) or not all(isinstance(x, str) for x in selected)
            or not set(selected) <= ids or not isinstance(summary, dict)
            or any(not isinstance(summary.get(k), str) or len(summary[k]) > 20000 for k, _ in SUMMARY_FIELDS)):
            return jsonify(error='Report content changed or settings are invalid. Reload the builder.'), 400
        for graph in case.get('graphs', []):
            include = 'graph:' + graph['id'] in selected
            if bool(graph.get('report', {}).get('include')) != include:
                graph.setdefault('report', {})['include'] = include
                graph['revision'] += 1
        case['report_builder'] = dict(order=order, selected=selected, summary={k: summary[k] for k, _ in SUMMARY_FIELDS})
        case['updated_at'] = app.utc_now()
        app.save_case(case)
    return jsonify(updated_at=case['updated_at'])


@bp.get('/cases/<case_id>/report-preview')
def preview(case_id):
    import app
    from graph_reports import legend_ink
    return render_template('report_preview.html', legend_ink=legend_ink, **compose(app.load_case(case_id)))


@bp.get('/cases/<case_id>/report-preview.pdf')
def preview_pdf(case_id):
    import app
    return Response(pdf(app.load_case(case_id)), mimetype='application/pdf',
                    headers={'Content-Disposition': 'inline; filename="report-preview.pdf"'})


def pdf(case):
    from pathlib import Path
    from xml.sax.saxutils import escape

    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer

    from graph_reports import legend_ink
    fonts = Path(reportlab.__file__).parent / 'fonts'
    for name, file in [('Builder', 'Vera.ttf'), ('BuilderBold', 'VeraBd.ttf')]:
        pdfmetrics.registerFont(TTFont(name, str(fonts / file)))
    styles = getSampleStyleSheet()
    for name in ('Normal', 'Title', 'Heading1', 'Heading2'):
        styles[name].fontName = 'BuilderBold' if name != 'Normal' else 'Builder'
        styles[name].textColor = colors.HexColor('#162938')
    styles['Normal'].fontSize = 10
    styles['Normal'].leading = 15
    styles['Normal'].spaceAfter = 4 if case.get('is_demo') else 8
    styles['Heading1'].keepWithNext = True
    styles['Heading2'].keepWithNext = True
    data = compose(case); story = []
    def add(text, style='Normal'):
        # Bound paragraphs so long captured findings remain splittable.
        for chunk in str(text).split('\n'):
            story.append(Paragraph(escape(chunk), styles[style]))
    add(data['display_name'], 'Title'); add('Investigation report | OSINT Menace')
    add('Executive summary', 'Heading1')
    for label, text in data['summary']:
        add(label, 'Heading2'); add(text or 'Not recorded.')
    for block in data['blocks']:
        add(block['label'], 'Heading1')
        if block.get('image'):
            image = Image(BytesIO(block['image']))
            ratio = min(6.9 * inch / image.imageWidth, 4 * inch / image.imageHeight)
            image.drawWidth, image.drawHeight = image.imageWidth * ratio, image.imageHeight * ratio
            story.append(image); story.append(Spacer(1, 10))
        add(block['text'])
        for entry in block.get('legend', []):
            story.append(Paragraph('<font color="' + legend_ink(entry['color']) + '" backColor="' + entry['color'] + '">'
                                   + escape(entry['type']) + '</font>  ' + escape(entry['description']), styles['Normal']))
        for label, records in block.get('details', []):
            if records:
                add(label, 'Heading2')
            for title, value in records:
                start = len(story)
                add(title); add(value)
                if len(title) + len(value) < 1800:
                    story[start:] = [KeepTogether(story[start:])]
    if data['appendix']:
        add('Evidence sources', 'Heading1')
        for source in data['appendix']:
            start = len(story)
            add(source['label'], 'Heading2'); add(source['text'])
            if len(source['text']) < 1800:
                story[start:] = [KeepTogether(story[start:])]
    out = BytesIO()
    def footer(canvas, doc):
        canvas.setFont('Builder', 8)
        canvas.drawString(54, 30, 'OSINT Menace | ' + ('Shareable report' if case.get('investigation', {}).get('redaction', {}).get('enabled') else str(case['id'])))
        canvas.drawRightString(558, 30, 'Page ' + str(doc.page))
    SimpleDocTemplate(out, pagesize=(612, 792), leftMargin=54, rightMargin=54,
                      topMargin=48, bottomMargin=54, title=data['display_name'], author='OSINT Menace').build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def word(case):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, RGBColor

    from graph_reports import legend_ink
    from word_report import clean_text
    data = compose(case); document = Document()
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    for border in document.styles['Title'].element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    document.add_paragraph(clean_text(data['display_name']), 'Title')
    document.add_paragraph('Investigation report | OSINT Menace')
    document.add_heading('Executive summary', 1)
    for label, text in data['summary']:
        document.add_heading(label, 2); document.add_paragraph(clean_text(text or 'Not recorded.'))
    for block in data['blocks']:
        document.add_heading(clean_text(block['label']), 1)
        if block.get('image'):
            from PIL import Image
            image = Image.open(BytesIO(block['image']))
            width = min(6, 4 * image.width / image.height)
            picture = document.add_picture(BytesIO(block['image']), width=Inches(width))
            picture._inline.docPr.set('descr', clean_text(block['label']))
        document.add_paragraph(clean_text(block['text']))
        for entry in block.get('legend', []):
            p = document.add_paragraph(); run = p.add_run(' ' + clean_text(entry['type']) + ' ')
            run.font.color.rgb = RGBColor.from_string(legend_ink(entry['color'])[1:])
            shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), entry['color'][1:])
            run._r.get_or_add_rPr().append(shading)
            p.add_run('  ' + clean_text(entry['description']))
        for label, records in block.get('details', []):
            if records:
                document.add_heading(label, 2)
            for title, value in records:
                p = document.add_paragraph(); p.add_run(clean_text(title)).bold = True
                document.add_paragraph(clean_text(value))
    if data['appendix']:
        document.add_heading('Evidence sources', 1)
        for source in data['appendix']:
            document.add_heading(clean_text(source['label']), 2); document.add_paragraph(clean_text(source['text']))
    out = BytesIO(); document.save(out); out.seek(0)
    return out


def markdown(case):
    data = compose(case)
    lines = ['# ' + data['display_name'], '', '## Executive summary', '']
    for label, text in data['summary']:
        lines.extend(['### ' + label, '', text or 'Not recorded.', ''])
    for block in data['blocks']:
        lines.extend(['## ' + block['label'], '', block['text'], ''])
        if block.get('image'):
            lines.append('[Graph or screenshot image is available in the PDF, Word and HTML reports.]')
        for entry in block.get('legend', []):
            lines.append(entry['type'] + ': ' + entry['description'])
        for label, records in block.get('details', []):
            lines.extend(['### ' + label, ''])
            for title, text in records:
                lines.extend([title, '', text, ''])
    if data['appendix']:
        lines.extend(['## Evidence sources', ''])
        for source in data['appendix']:
            lines.extend(['### ' + source['label'], '', source['text'], ''])
    return '\n'.join(lines)
